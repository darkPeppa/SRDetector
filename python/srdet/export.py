"""Readable exports of runs, studies and optimisations.

    srdet export runs/<run>                 -> runs/<run>/export/<name>.xlsx
    srdet export runs/<run> --format zip    -> runs/<run>/export/<name>.zip

Excel: styled sheets (summary with Russian labels, parameters, events, histograms, plots
embedded as pictures). ZIP: the Excel file plus CSV / Parquet tables, plots, configuration,
the ROOT file and a README describing every column.
"""
from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pandas as pd
import yaml

MAX_ROWS = 200_000  # Excel sheet limit is ~1M rows; keep files light

COLUMN_DOCS = {
    "evt": "номер события", "e0_keV": "энергия первичной частицы, кэВ", "pdg": "PDG-код первичной частицы",
    "x0_mm": "x точки старта, мм", "y0_mm": "y точки старта, мм",
    "first_proc": "первое взаимодействие: 0 нет, 1 фото, 2 комптон, 3 рэлей, 4 пары, 5 другое",
    "first_z_mm": "глубина первого взаимодействия, мм", "first_kind": "объём первого взаимодействия (код)",
    "first_cell": "сота первого взаимодействия", "edep_total_keV": "энергия во всём детекторе, кэВ",
    "edep_core_keV": "энергия в сцинтилляторе, кэВ", "edep_wall_keV": "энергия в стенках, кэВ",
    "edep_coating_keV": "энергия в покрытии, кэВ", "edep_matrix_keV": "энергия в ободе матрицы, кэВ",
    "edep_window_keV": "энергия во входном окне, кэВ", "edep_housing_keV": "энергия в корпусе, кэВ",
    "edep_coupling_keV": "энергия в смазке, кэВ", "edep_sensor_keV": "энергия в сенсоре, кэВ",
    "edep_slab_keV": "энергия в блоке (режим slab), кэВ", "edep_world_keV": "энергия в воздухе мира, кэВ",
    "e_escape_keV": "энергия, унесённая из мира, кэВ", "n_cells_hit": "сот с отложением выше порога",
    "edep_first_cell_keV": "энергия в соте первого взаимодействия, кэВ", "edep_max_cell_keV": "максимум по сотам, кэВ",
    "n_phot": "фотоэффектов", "n_compt": "комптоновских рассеяний", "n_rayl": "рэлеевских рассеяний",
    "n_conv": "рождений пар", "n_eioni": "ионизаций электронами (δ)", "n_ebrem": "тормозных квантов",
    "n_annihil": "аннигиляций", "n_msc": "шагов многократного рассеяния", "n_fluo": "флуоресцентных квантов",
    "n_fluo_escape": "флуоресцентных квантов, покинувших детектор", "n_scint": "рождено сцинтилляционных фотонов",
    "n_cher": "рождено черенковских фотонов", "n_cher_primary": "черенковских от первичной частицы",
    "op_fresnel_refl": "френелевских отражений", "op_tir": "полных внутренних отражений",
    "op_lambert": "ламбертовских отражений", "op_spike_lobe": "зеркальных отражений (spike/lobe)",
    "op_lut_refl": "отражений по LUT", "op_refraction": "преломлений/прохождений", "op_surf_abs": "поглощений на поверхностях",
    "op_detection": "детектирований", "op_norindex": "уходов в неоптический материал",
    "n_det": "детектировано фотонов", "n_abs_bulk": "поглощено в объёме", "n_abs_surface": "поглощено на поверхностях",
    "n_norindex": "ушло в неоптический материал", "n_escaped": "покинуло мир", "n_lost_other": "прочие потери",
    "abs_bulk_core": "объёмное поглощение в сцинтилляторе", "abs_bulk_wall": "объёмное поглощение в стенках",
    "abs_bulk_coupling": "объёмное поглощение в смазке", "abs_bulk_other": "объёмное поглощение в прочих",
    "coat_refl": "отражений от стенки канала", "coat_abs": "поглощений стенкой канала",
    "coat_trans": "прохождений сквозь стенку канала", "n_exit": "фотонов, вышедших из каналов к сенсору",
    "det_same_cell": "детектировано из своей соты", "det_other_cell": "детектировано из чужой соты",
    "n_op_rayleigh": "рэлеевских рассеяний оптических фотонов", "first_boundary_status": "статус первой границы (тесты)",
    "x_mm": "x на сенсоре, мм", "y_mm": "y на сенсоре, мм", "t_ns": "время прихода, нс", "lambda_nm": "длина волны, нм",
    "n_refl": "отражений", "birth_cell": "сота рождения", "exit_cell": "сота выхода", "birth_z_mm": "глубина рождения, мм",
    "n_coat_refl": "отражений от стенки канала", "cell": "номер соты", "edep_keV": "энергия, кэВ",
}


def _labels():
    from .labels import METRIC_LABELS
    return METRIC_LABELS


def _flat(d: dict, prefix: str = "") -> list[tuple[str, object]]:
    out = []
    for k, v in d.items():
        if str(k).startswith("_"):
            continue
        p = f"{prefix}{k}"
        if isinstance(v, dict):
            out += _flat(v, p + ".")
        else:
            out.append((p, v if not isinstance(v, list) else json.dumps(v, ensure_ascii=False)))
    return out


def _json(p: Path) -> dict:
    return json.loads(p.read_text()) if p.exists() else {}


# ------------------------------------------------------------------ styling helpers
def _style(ws, header_row: int = 1, widths: dict | None = None, freeze: str | None = "A2"):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    fill = PatternFill("solid", fgColor="2A6FDB")
    for c in ws[header_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for col in ws.iter_cols(min_row=header_row, max_row=min(ws.max_row, header_row + 200)):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, width + 2), 60)
    for k, w in (widths or {}).items():
        ws.column_dimensions[k].width = w
    if freeze:
        ws.freeze_panes = freeze


def _df_sheet(wb, title: str, df: pd.DataFrame, docs: bool = True):
    ws = wb.create_sheet(title[:31])
    cols = list(df.columns)
    ws.append(cols)
    if docs:  # second row: human description of the column
        ws.append([COLUMN_DOCS.get(c, "") for c in cols])
    for row in df.head(MAX_ROWS).itertuples(index=False):
        ws.append([None if (isinstance(v, float) and v != v) else v for v in row])
    _style(ws, freeze="A3" if docs else "A2")
    if docs:
        from openpyxl.styles import Font
        for c in ws[2]:
            c.font = Font(italic=True, color="667085", size=9)
    if len(df) > MAX_ROWS:
        ws.cell(row=ws.max_row + 2, column=1, value=f"показаны первые {MAX_ROWS} строк из {len(df)} (полностью — в CSV/Parquet)")
    return ws


def _plots_sheet(wb, plots: list[Path], title: str = "Графики"):
    if not plots:
        return
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font
    ws = wb.create_sheet(title)
    ws["A1"] = "Графики (кликните по картинке, чтобы выделить; оригиналы PNG — в ZIP-пакете)"
    ws["A1"].font = Font(bold=True, size=12)
    row = 3
    for i, p in enumerate(plots):
        col = "A" if i % 2 == 0 else "N"
        ws[f"{col}{row}"] = p.stem.replace("_", " ")
        ws[f"{col}{row}"].font = Font(bold=True)
        img = XLImage(str(p))
        scale = 620 / max(img.width, 1)
        img.width, img.height = img.width * scale, img.height * scale
        ws.add_image(img, f"{col}{row + 1}")
        if i % 2 == 1:
            row += int(img.height / 19) + 4
    return ws


def _summary_sheet(wb, title: str, rows: list[tuple]):
    from openpyxl.styles import Font
    ws = wb.active
    ws.title = "Сводка"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    r = 3
    for block, items in rows:
        ws.cell(row=r, column=1, value=block).font = Font(bold=True, size=12, color="2A6FDB")
        r += 1
        for item in items:
            for j, v in enumerate(item, start=1):
                ws.cell(row=r, column=j, value=v)
            r += 1
        r += 1
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 14
    return ws


# ------------------------------------------------------------------ runs
def _run_tables(d: Path) -> dict[str, pd.DataFrame]:
    out = {}
    if (d / "srdet.root").exists():
        import uproot
        f = uproot.open(d / "srdet.root")
        for name in ("events", "photons", "cells"):
            if name in f and f[name].num_entries > 0:
                out[name] = f[name].arrays(library="pd")
        hist = []
        for key, cls in f.classnames().items():
            if cls.startswith("TH1"):
                h = f[key]
                v, e = h.to_numpy()
                hist.append(pd.DataFrame({f"{key.split(';')[0]} | от": e[:-1], f"{key.split(';')[0]} | до": e[1:],
                                          f"{key.split(';')[0]} | число": v}))
        if hist:
            out["histograms"] = pd.concat(hist, axis=1)
    return out


def run_xlsx(d: Path, out: Path) -> Path:
    from openpyxl import Workbook
    labels = _labels()
    m, meta = _json(d / "metrics.json"), _json(d / "meta.json")
    cfg = yaml.safe_load((d / "config.yaml").read_text()) if (d / "config.yaml").exists() else {}
    wb = Workbook()
    from .labels import KEY_ORDER
    keys = [k for k in KEY_ORDER if k in m] + [k for k in m if k not in KEY_ORDER]
    metric_rows = [(labels.get(k, k), k, m[k], m.get(f"{k}_err")) for k in keys
                   if not k.endswith("_err") and k not in ("name", "sanity_problems")
                   and not isinstance(m[k], (list, dict))]
    g, b = cfg.get("geometry", {}), cfg.get("beam", {})
    _summary_sheet(wb, f"Прогон {meta.get('name', d.name)}", [
        ("Прогон", [("статус", (d / "STATUS").read_text().strip() if (d / "STATUS").exists() else ""),
                    ("начат", meta.get("started")), ("время счёта, с", meta.get("wall_time_s")),
                    ("Geant4", meta.get("geant4")), ("git commit", meta.get("git_commit"))]),
        ("Детектор и пучок", [("режим", g.get("mode")), ("длина канала", g.get("length")), ("шаг", g.get("pitch")),
                              ("диаметр канала", g.get("channel_diameter")), ("сцинтиллятор", g.get("core_material")),
                              ("покрытие", cfg.get("optics", {}).get("coating")), ("сенсор", g.get("sensor")),
                              ("пучок", f"{b.get('particle')} {b.get('energy')} {b.get('type')}")]),
        ("Метрики (название · ключ · значение · ошибка)", metric_rows),
    ])
    ws = wb.create_sheet("Параметры")
    ws.append(["параметр", "значение"])
    for k, v in _flat(cfg):
        ws.append([k, str(v) if not isinstance(v, (int, float, bool)) else v])
    _style(ws)
    tables = _run_tables(d)
    titles = {"events": "События", "photons": "Фотоны", "cells": "Соты", "histograms": "Гистограммы"}
    for k, df in tables.items():
        _df_sheet(wb, titles[k], df, docs=k != "histograms")
    _plots_sheet(wb, sorted((d / "plots").glob("*.png")) + sorted(d.glob("*.png")))
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# ------------------------------------------------------------------ studies / optimisations
def table_xlsx(d: Path, out: Path) -> Path:
    from openpyxl import Workbook
    labels = _labels()
    wb = Workbook()
    is_opt = (d / "trials.csv").exists()
    spec = yaml.safe_load(((d / "spec.yaml") if is_opt else (d / "study.yaml")).read_text()) \
        if ((d / "spec.yaml").exists() or (d / "study.yaml").exists()) else {}
    items = [("название", spec.get("name")), ("заголовок", spec.get("title")),
             ("базовая конфигурация", spec.get("base"))]
    if is_opt:
        items += [("цель", spec.get("objective", {}).get("expr")),
                  ("ограничения", "; ".join(spec.get("constraints", [])))]
        final = d / "final.csv"
        if final.exists():
            best = pd.read_csv(final).iloc[0]
            items += [(f"лучшее: {a}", best.get(a)) for a in spec.get("space", {})]
            items += [("лучшее: цель", best.get("objective")), ("± σ", best.get("objective_std"))]
    else:
        items += [("сканируемые параметры", ", ".join(spec.get("scan", {}) or ["варианты"]))]
    _summary_sheet(wb, spec.get("title", d.name), [("Описание", items),
                                                    ("Что в листах", [("Результаты / Испытания", "все точки"),
                                                                      ("Лучшие", "проверенные кандидаты (оптимизация)"),
                                                                      ("Графики", "картинки отчёта")])])
    if is_opt:
        _df_sheet(wb, "Испытания", pd.read_csv(d / "trials.csv"), docs=False)
        if (d / "final.csv").exists():
            _df_sheet(wb, "Лучшие", pd.read_csv(d / "final.csv"), docs=False)
    elif (d / "results.csv").exists():
        df = pd.read_csv(d / "results.csv")
        df = df.rename(columns={c: f"{labels[c]} ({c})" for c in df.columns if c in labels})
        _df_sheet(wb, "Результаты", df, docs=False)
    _plots_sheet(wb, sorted((d / "plots").glob("*.png")))
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# ------------------------------------------------------------------ zip bundle
README_HEAD = """# Экспорт SRDetector: {name}

* `{name}.xlsx` — то же в Excel: сводка, параметры, таблицы, графики
* `csv/` и `parquet/` — таблицы целиком (Parquet читается pandas / polars / R arrow)
* `plots/` — графики PNG, `config/` — конфигурация и макрос Geant4
* `srdet.root` (если есть) — гистограммы и ntuple для ROOT / JSROOT / uproot

"""


def bundle(d: Path, out: Path) -> Path:
    name = d.name
    xlsx = out.with_suffix(".xlsx")
    (table_xlsx if (d / "trials.csv").exists() or (d / "results.csv").exists() else run_xlsx)(d, xlsx)
    readme = README_HEAD.format(name=name)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, f"{name}/{xlsx.name}")
        for p in sorted(d.glob("plots/*.png")) + sorted(d.glob("*.png")):
            z.write(p, f"{name}/plots/{p.name}")
        for f in ("config.yaml", "run.mac", "meta.json", "metrics.json", "geometry.json", "materials.json",
                  "spec.yaml", "study.yaml", "best.yaml", "report.md", "results.md", "STATUS"):
            if (d / f).exists():
                z.write(d / f, f"{name}/config/{f}" if f.endswith((".yaml", ".mac", ".json", "STATUS")) else f"{name}/{f}")
        for f in ("trials.csv", "final.csv", "results.csv"):
            if (d / f).exists():
                z.write(d / f, f"{name}/csv/{f}")
        if (d / "srdet.root").exists():
            z.write(d / "srdet.root", f"{name}/srdet.root")
            tables = _run_tables(d)
            for k, df in tables.items():
                z.writestr(f"{name}/csv/{k}.csv", df.to_csv(index=False))
                if k != "histograms":
                    tmp = out.parent / f".{k}.parquet"
                    df.to_parquet(tmp, compression="zstd")
                    z.write(tmp, f"{name}/parquet/{k}.parquet")
                    tmp.unlink()
            if "events" in tables:
                readme += "## Столбцы таблицы событий (`csv/events.csv`)\n\n| столбец | смысл |\n|---|---|\n"
                readme += "".join(f"| `{c}` | {COLUMN_DOCS.get(c, '')} |\n" for c in tables["events"].columns)
        z.writestr(f"{name}/README.md", readme)
    return out


def export(path: str | Path, fmt: str = "xlsx") -> Path:
    d = Path(path).resolve()
    out = d / "export" / f"{d.name}.{fmt}"
    if fmt == "zip":
        return bundle(d, out)
    if (d / "trials.csv").exists() or (d / "results.csv").exists():
        return table_xlsx(d, out)
    return run_xlsx(d, out)
