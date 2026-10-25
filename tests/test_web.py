"""Web API tests (FastAPI TestClient, no Geant4): schema, configs, macro preview, run browsing,
geometry, file access control, job validation and result parsing."""
from __future__ import annotations

import json
import shutil
import uuid

import pytest

from conftest import ROOT

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient  # noqa: E402

from srdet.web.jobs import JobManager  # noqa: E402
from srdet.web.server import create_app  # noqa: E402


@pytest.fixture(scope="module", autouse=True)
def repo_runs():
    """The physics tests point SRDET_RUNS to a temporary folder; the web API serves runs/ of the repo."""
    import os
    old = os.environ.get("SRDET_RUNS")
    os.environ["SRDET_RUNS"] = str(ROOT / "runs")
    yield
    if old is None:
        os.environ.pop("SRDET_RUNS", None)
    else:
        os.environ["SRDET_RUNS"] = old


@pytest.fixture(scope="module")
def fake_run(repo_runs):
    d = ROOT / "runs" / f"zz_test_web_{uuid.uuid4().hex[:6]}"
    (d / "plots").mkdir(parents=True)
    (d / "config.yaml").write_text("name: fake\ngeometry: {mode: unit, length: 2 mm, coating_thickness: 2 um}\n"
                                   "optics: {coating: Ag}\nbeam: {particle: gamma, energy: 60 keV, type: pencil}\n")
    (d / "metrics.json").write_text(json.dumps({"det_yield_per_keV": 8.2, "detection_efficiency": 0.9,
                                                "lce": float("nan"), "sanity_problems": []}))
    (d / "meta.json").write_text(json.dumps({"name": "fake", "wall_time_s": 1.5}))
    (d / "STATUS").write_text("ok\n")
    (d / "geometry.json").write_text(json.dumps({"mode": "unit", "pitch_mm": 0.208, "channel_diameter_mm": 0.2,
                                                 "matrix_radius_mm": 0.121, "cells": [[0, 0, 0.0, 0.0]],
                                                 "n_cells": 1, "length_mm": 2, "z_front_mm": -1.5,
                                                 "z_back_mm": 1.07, "z_sensor_mm": 1.05}))
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(2, 1.5))
    ax.plot([0, 1], [0, 1])
    fig.savefig(d / "plots" / "edep_spectrum.png")
    plt.close(fig)
    import numpy as np
    import uproot
    with uproot.recreate(d / "srdet.root") as f:  # small ROOT file: one histogram + an events tree
        f["n_det"] = (np.array([0.0, 5, 9, 3]), np.array([0.0, 10, 20, 30, 40]))
        f["events"] = {"evt": np.arange(4, dtype=np.int32), "n_det": np.array([0, 12, 25, 31], dtype=np.int32)}
    yield d.relative_to(ROOT).as_posix()
    shutil.rmtree(d, ignore_errors=True)


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    jm = JobManager(tmp_path_factory.mktemp("jobs"))
    return TestClient(create_app(jm))


def test_index_and_static(client):
    assert "SRDetector" in client.get("/").text
    assert client.get("/static/app.js").status_code == 200
    assert client.get("/static/app.js").headers["cache-control"] == "no-cache"


def test_schema_has_typed_fields(client):
    s = client.get("/api/schema").json()
    fields = {f["path"]: f for sec in s["sections"] for f in sec["fields"]}
    assert fields["geometry.length"]["kind"] == "quantity" and fields["geometry.length"]["unit_kind"] == "length"
    assert "Ag" in fields["optics.coating"]["choices"]
    assert fields["geometry.check_overlaps"]["kind"] == "bool"
    assert fields["beam.spectrum"]["kind"] == "special"
    assert fields["geometry.pitch"]["help"]  # trailing comment of defaults.yaml
    assert "Am241" in s["spectra"]["sources"]


def test_configs_and_macro(client):
    names = [c["name"] for c in client.get("/api/configs").json()]
    assert "full" in names and "unit_optical" in names
    cfg = client.get("/api/configs/full").json()
    assert cfg["geometry"]["mode"] == "full"
    r = client.post("/api/macro", json={"base": "unit_xray", "overrides": {"geometry.length": "2 mm"}})
    assert r.status_code == 200 and "/srdet/geom/length 2 mm" in r.json()["macro"]


def test_run_browsing(client, fake_run):
    runs = client.get("/api/runs").json()
    me = next(r for r in runs if r["path"] == fake_run)
    assert me["metrics"]["det_yield_per_keV"] == 8.2 and me["coating"] == "Ag"
    run = client.get("/api/run", params={"path": fake_run}).json()
    assert run["metrics_all"]["lce"] is None          # NaN is sent as null
    assert run["plots"][0]["path"].endswith("plots/edep_spectrum.png")
    assert run["config"]["optics"]["coating"] == "Ag"
    g = client.get("/api/run/geometry", params={"path": fake_run}).json()
    assert g["n_cells"] == 1 and g["config_optics"]["coating"] == "Ag"
    assert client.get(f"/files/{fake_run}/plots/edep_spectrum.png").status_code == 200


def test_file_access_is_restricted(client):
    assert client.get("/files/python/srdet/web/server.py").status_code == 403
    assert client.get("/files/runs/../../etc/passwd").status_code in (403, 404)
    assert client.get("/api/doc", params={"path": "app/srdet.cc"}).status_code == 403
    assert client.get("/api/doc", params={"path": "README.md"}).json()["markdown"].startswith("#")


def test_job_validation(client):
    assert client.post("/api/jobs", json={"kind": "study", "target": "../evil"}).status_code == 400
    assert client.post("/api/jobs", json={"kind": "rm -rf"}).status_code == 400


def test_job_result_parsing():
    log = "[srdet] run x: /work/runs/a\n[srdet] study U01 done: /work/runs/studies/20261024_U01\n"
    assert JobManager._result(log).endswith("runs/studies/20261024_U01")
    assert JobManager._result("[srdet] optimize done: /work/runs/optimize/O1/report.md").endswith("runs/optimize/O1")


def test_studies_and_optimizations_listing(client):
    st = client.get("/api/studies").json()
    assert any(s["name"] == "U01_coating" for s in st["specs"])
    op = client.get("/api/optimizations").json()
    assert any(s["name"].startswith("O01") for s in op["specs"])


class DummyJobs:
    """Records submissions instead of running srdet in a subprocess."""
    def __init__(self):
        self.submitted = []

    def submit(self, kind, args, title, meta=None):
        job = {"id": f"j{len(self.submitted)}", "kind": kind, "args": args, "title": title, "meta": meta or {},
               "status": "queued", "created": "2026-01-01T00:00:00", "result": None}
        self.submitted.append(job)
        return job

    def get(self, jid, tail=200):
        return next((dict(j, log_tail="") for j in self.submitted if j["id"] == jid), None)

    def list(self):
        return self.submitted

    def cancel(self, jid):
        return False


def test_tracks_and_scan_jobs(fake_run):
    jobs = DummyJobs()
    c = TestClient(create_app(jobs))
    r = c.post("/api/jobs", json={"kind": "tracks", "config_path": f"{fake_run}/config.yaml", "events": 7}).json()
    assert "run.tracks=7" in r["args"] and r["meta"]["source_run"] == fake_run
    r = c.post("/api/jobs", json={"kind": "scan", "path": fake_run, "alias": "L", "points": 4}).json()
    assert r["args"][:2] == ["optimize", "--resume"] and "--scan" in r["args"]
    assert c.post("/api/jobs", json={"kind": "scan", "path": fake_run, "alias": "1; rm"}).status_code == 400


def test_histogram_api(client, fake_run):
    h = client.get("/api/run/hist", params={"path": fake_run, "name": "n_det"}).json()
    assert h["values"] == [0, 5, 9, 3] and h["edges"][0][-1] == 40 and h["entries"] == 17
    assert client.get("/api/run/hist", params={"path": fake_run, "name": "nope"}).status_code == 404


def test_export_xlsx_and_zip(client, fake_run):
    import io
    import zipfile
    import openpyxl
    r = client.get("/api/export", params={"path": fake_run, "fmt": "xlsx"})
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert {"Сводка", "Параметры", "События", "Гистограммы", "Графики"} <= set(wb.sheetnames)
    labels = [row[0] for row in wb["Сводка"].iter_rows(values_only=True) if row[0]]
    assert "Фотоэлектронов на кэВ" in labels
    z = zipfile.ZipFile(io.BytesIO(client.get("/api/export", params={"path": fake_run, "fmt": "zip"}).content))
    names = z.namelist()
    assert any(n.endswith("README.md") for n in names) and any(n.endswith("csv/events.csv") for n in names)
    assert "n_det" in z.read(next(n for n in names if n.endswith("README.md"))).decode()


def test_skins():
    for skin in ("pink", "rap"):
        c = TestClient(create_app(DummyJobs(), skin=skin))
        html = c.get("/").text
        assert f"/static/themes/{skin}.css" in html and "/static/themes/skin.js" in html
        assert c.get("/api/skin").json()["name"] == skin
        assert c.get(f"/static/themes/{skin}.js").status_code == 200
    with pytest.raises(ValueError):
        create_app(DummyJobs(), skin="neon")
    assert "themes" not in TestClient(create_app(DummyJobs())).get("/").text
